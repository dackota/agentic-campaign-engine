"""What each preview route renders from the scripted agent's output matches a snapshot. HTML
and SVG snapshot as text; PDF and PPTX as the text they hold. The routes take their renderer
and the steps it reads from `_ASSET_KINDS`, so this also pins which steps feed each kind."""

import io
import re
from collections.abc import Callable
from pathlib import Path

import pytest
from campaign_api.main import _ASSET_KINDS
from pptx import Presentation
from pypdf import PdfReader
from render_service.assets import render_graphic_svg, render_landing_page_html
from render_service.email import MAX_EMAIL_BYTES, render_email

from tests.support.scripted_outputs import scripted_output

SNAPSHOTS = Path(__file__).parent / "snapshots"


def _pdf_text(content: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)


def _pptx_text(content: bytes) -> str:
    deck = Presentation(io.BytesIO(content))
    return "\n---\n".join(
        "\n".join(s.text_frame.text.replace("\v", "\n") for s in slide.shapes if s.has_text_frame)
        for slide in deck.slides
    )


def _text(content: str | bytes) -> str:
    assert isinstance(content, str)
    return content


# kind -> (snapshot file, how to read the rendered content as text)
SNAPSHOT_KINDS: dict[str, tuple[str, Callable[..., str]]] = {
    "one-pager": ("one_pager_preview.txt", _pdf_text),
    "deck": ("deck_preview.txt", _pptx_text),
    "landing-page": ("landing_page_preview.html", _text),
    "graphic": ("graphic_preview.svg", _text),
}


def _rendered(kind: str) -> str | bytes:
    steps, render, _ = _ASSET_KINDS[kind]
    return render(*(scripted_output(step) for step in steps))


@pytest.mark.parametrize("kind", SNAPSHOT_KINDS)
def test_preview_matches_snapshot(kind: str) -> None:
    snapshot, as_text = SNAPSHOT_KINDS[kind]

    assert as_text(_rendered(kind)) == (SNAPSHOTS / snapshot).read_text()


def test_landing_page_preview_is_the_same_bytes_every_time() -> None:
    page = scripted_output("landing_page")

    assert render_landing_page_html(page) == render_landing_page_html(page)


def test_graphic_preview_is_the_same_bytes_every_time() -> None:
    one_pager = scripted_output("one_pager")

    assert render_graphic_svg(one_pager) == render_graphic_svg(one_pager)


def test_email_preview_matches_snapshot() -> None:
    draft = scripted_output("email_writer")

    assert render_email(draft) == (SNAPSHOTS / "email_preview.html").read_text()


def test_email_preview_keeps_the_logo_inline() -> None:
    html = render_email(scripted_output("email_writer"))

    assert "assets.example.com" not in html
    assert 'src="data:image/' in html
    assert len(html.encode()) < MAX_EMAIL_BYTES
    assert all(re.search(r'alt="[^"]+"', tag) for tag in re.findall(r"<img\b[^>]*>", html))
