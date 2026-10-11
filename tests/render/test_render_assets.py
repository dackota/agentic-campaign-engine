"""The asset renderers: one-pager PDF, deck PPTX, landing page HTML, graphic SVG."""

import io
from xml.etree import ElementTree

import pytest
from pptx import Presentation
from pypdf import PdfReader
from render_service.assets import (
    render_deck_pptx,
    render_graphic_svg,
    render_landing_page_html,
    render_one_pager_pdf,
)

ONE_PAGER = {
    "title": "Volt Typhoon",
    "summary": "Edge routers are hit.",
    "bullets": ["Patch", "Rotate"],
}
LANDING = {
    "headline": "Volt Typhoon",
    "subheadline": "What it means",
    "body": "Line one.\n\nLine two.",
    "call_to_action": "Read more",
}
HOSTILE = '<script>alert(1)</script> & "quotes" \' <img src=x onerror=y>'


def test_one_pager_pdf_has_text_and_is_pdf() -> None:
    pdf = render_one_pager_pdf(ONE_PAGER)

    assert pdf.startswith(b"%PDF")
    pages = PdfReader(io.BytesIO(pdf)).pages
    assert len(pages) == 1
    text = pages[0].extract_text()
    assert "Volt Typhoon" in text and "Patch" in text and "Rotate" in text


def test_pdf_is_deterministic() -> None:
    assert render_one_pager_pdf(ONE_PAGER) == render_one_pager_pdf(ONE_PAGER)


def test_deck_pptx_has_slides() -> None:
    deck = Presentation(io.BytesIO(render_deck_pptx(ONE_PAGER, LANDING)))

    texts = [
        shape.text_frame.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert len(deck.slides) >= 3
    assert "Volt Typhoon" in texts[0]
    assert any("Patch" in t for t in texts)
    assert any("Read more" in t for t in texts)


def test_deck_is_deterministic() -> None:
    assert render_deck_pptx(ONE_PAGER, LANDING) == render_deck_pptx(ONE_PAGER, LANDING)


def test_landing_page_html_has_content() -> None:
    html = render_landing_page_html(LANDING)

    assert html.startswith("<!doctype html>")
    assert "<h1>Volt Typhoon</h1>" in html
    assert "Line two." in html and "Read more" in html


@pytest.mark.parametrize("hostile", [HOSTILE, "", "x" * 50_000, "\x00\x07 ünïcode ☃"])
def test_landing_page_html_escapes_hostile_text(hostile: str) -> None:
    html = render_landing_page_html({**LANDING, "headline": hostile or "h", "body": hostile or "b"})

    assert "<script>" not in html
    assert "<img" not in html


@pytest.mark.parametrize(
    "hostile", [HOSTILE, "x" * 5_000, "\x00\x07 ünïcode ☃", "a" * 3 + "\n" * 5]
)
def test_graphic_svg_is_well_formed_xml(hostile: str) -> None:
    svg = render_graphic_svg({**ONE_PAGER, "title": hostile})

    root = ElementTree.fromstring(svg)
    assert root.tag.endswith("svg")
    assert "<script" not in svg
    assert "<img" not in svg


@pytest.mark.parametrize("hostile", [HOSTILE, "x" * 5_000, "\x00 ünïcode ☃"])
def test_pdf_and_deck_survive_hostile_text(hostile: str) -> None:
    pager = {"title": hostile, "summary": hostile, "bullets": [hostile]}

    assert render_one_pager_pdf(pager).startswith(b"%PDF")
    assert (
        len(Presentation(io.BytesIO(render_deck_pptx(pager, {**LANDING, "body": hostile}))).slides)
        >= 3
    )
