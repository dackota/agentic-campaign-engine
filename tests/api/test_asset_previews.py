"""Seam 1 at fast depth: each asset kind has a preview route, and its output matches a
snapshot. HTML and SVG snapshot as text; PDF and PPTX as the text they hold."""

import io
import uuid
from collections.abc import Callable
from pathlib import Path
from typing import Any

import httpx
import pytest
from ace_engine.default_workflow import DEFAULT_WORKFLOW_NAME
from pptx import Presentation
from pypdf import PdfReader

from tests.api.helpers import AWAITING, create_on, until_status
from tests.harness.stack import Stack

SNAPSHOTS = Path(__file__).parent / "snapshots"
REPORT = "Volt Typhoon targets edge routers\nActors exploit unpatched VPN appliances."


def _pdf_text(content: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)


def _pptx_text(content: bytes) -> str:
    deck = Presentation(io.BytesIO(content))
    return "\n---\n".join(
        "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        for slide in deck.slides
    )


# kind -> (content type, snapshot file, how to read the body as text)
KINDS: dict[str, tuple[str, str, Callable[[httpx.Response], str]]] = {
    "one-pager": ("application/pdf", "one_pager_preview.txt", lambda r: _pdf_text(r.content)),
    "deck": (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "deck_preview.txt",
        lambda r: _pptx_text(r.content),
    ),
    "landing-page": ("text/html", "landing_page_preview.html", lambda r: r.text),
    "graphic": ("image/svg+xml", "graphic_preview.svg", lambda r: r.text),
}


@pytest.fixture
def campaign(api: httpx.Client) -> dict[str, Any]:
    created = create_on(api, DEFAULT_WORKFLOW_NAME, REPORT)
    return until_status(api, created["id"], AWAITING)


def _preview(api: httpx.Client, campaign_id: str, kind: str) -> httpx.Response:
    return api.get(f"/campaigns/{campaign_id}/preview/{kind}")


@pytest.mark.parametrize("kind", KINDS)
def test_preview_returns_right_content_type(
    api: httpx.Client, campaign: dict[str, Any], kind: str
) -> None:
    response = _preview(api, campaign["id"], kind)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(KINDS[kind][0])
    assert response.content


@pytest.mark.parametrize("kind", KINDS)
def test_preview_matches_snapshot(api: httpx.Client, campaign: dict[str, Any], kind: str) -> None:
    _, snapshot, as_text = KINDS[kind]

    assert as_text(_preview(api, campaign["id"], kind)) == (SNAPSHOTS / snapshot).read_text()


@pytest.mark.parametrize("kind", KINDS)
def test_second_get_returns_same_bytes(
    api: httpx.Client, campaign: dict[str, Any], kind: str
) -> None:
    assert (
        _preview(api, campaign["id"], kind).content == _preview(api, campaign["id"], kind).content
    )


@pytest.mark.parametrize("kind", KINDS)
def test_preview_404_unknown_campaign(api: httpx.Client, kind: str) -> None:
    assert _preview(api, str(uuid.uuid4()), kind).status_code == 404


def test_unknown_kind_404(api: httpx.Client, campaign: dict[str, Any]) -> None:
    assert _preview(api, campaign["id"], "poster").status_code == 404


def test_preview_409_when_run_has_no_output(stack: Stack, api: httpx.Client) -> None:
    created = create_on(api, _no_writer_workflow(api))
    until_status(api, created["id"], AWAITING)

    assert _preview(api, created["id"], "graphic").status_code == 409


def _no_writer_workflow(api: httpx.Client) -> str:
    from tests.api.helpers import DRAFT_APPROVE_PUBLISH, save_workflow, workflow_name

    name = workflow_name()
    save_workflow(api, name, DRAFT_APPROVE_PUBLISH)
    return name
