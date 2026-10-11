"""Brand renderers for the other asset kinds: the one-pager as PDF, the sales deck as PPTX,
the landing page as HTML, and a graphic from an SVG template.

Pure and deterministic: the same content gives the same bytes. Text from an agent is
untrusted, so each renderer strips control characters and escapes what it writes.
"""

import io
import re
import zipfile
from collections.abc import Mapping, Sequence
from datetime import datetime
from html import escape
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer

BRAND_NAME = "Agentic Campaign Engine"
BRAND_COLOR = "#0b5fff"
BRAND_RGB = (0x0B, 0x5F, 0xFF)
TEXT_COLOR = "#1f2933"
FONT_STACK = "Helvetica, Arial, sans-serif"
PAGE_MAX_WIDTH = "640px"
FIXED_TIMESTAMP = datetime(2026, 1, 1)  # embedded dates would make equal content differ
ZIP_EPOCH = (2026, 1, 1, 0, 0, 0)
PDF_TITLE_SIZE = 22
PDF_TEXT_SIZE = 11
PDF_LEADING = 15
PDF_SPACE = 10
DECK_TITLE_PT = 36
DECK_TEXT_PT = 20
SVG_WIDTH = 1200
SVG_HEIGHT = 630
SVG_TITLE_MAX_CHARS = 60
SVG_TITLE_SIZE = 56
SVG_BRAND_SIZE = 28
SVG_MARGIN = 80
SVG_TITLE_Y = 330
SVG_BRAND_Y = 560
MAX_BULLETS = 20

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
_PARAGRAPH_BREAK = re.compile(r"\n\s*\n")

_LANDING_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{headline}</title>
<style>
body {{ margin: 0; font-family: {font}; color: {text}; background: #f4f6f8; }}
header {{ background: {color}; color: #fff; padding: 16px 24px; font-weight: bold; }}
main {{ max-width: {width}; margin: 24px auto; padding: 24px; background: #fff; }}
a.cta {{ display: inline-block; padding: 12px 24px; background: {color}; color: #fff;
  text-decoration: none; border-radius: 4px; }}
</style>
</head>
<body>
<header>{brand}</header>
<main>
<h1>{headline}</h1>
<h2>{subheadline}</h2>
{paragraphs}
<p><a class="cta" href="#act">{cta}</a></p>
</main>
</body>
</html>
"""

_SVG_TEMPLATE = """<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" \
viewBox="0 0 {w} {h}" role="img" aria-label="{title}">
  <title>{title}</title>
  <rect width="{w}" height="{h}" fill="{color}"/>
  <text x="{margin}" y="{title_y}" font-family="{font}" font-size="{title_size}" \
font-weight="bold" fill="#ffffff">{title}</text>
  <text x="{margin}" y="{brand_y}" font-family="{font}" font-size="{brand_size}" \
fill="#ffffff">{brand}</text>
</svg>
"""


def _clean(text: str) -> str:
    """Drop the characters XML cannot carry; keep everything else."""
    return _CONTROL_CHARS.sub("", text)


def _bullets(one_pager: Mapping[str, Any]) -> list[str]:
    bullets: Sequence[str] = one_pager["bullets"]
    return [_clean(b) for b in bullets[:MAX_BULLETS]]


def render_one_pager_pdf(one_pager: Mapping[str, Any]) -> bytes:
    """Render the one-pager (title, summary, bullets) as a PDF."""
    title_style = ParagraphStyle(
        "title",
        fontName="Helvetica-Bold",
        fontSize=PDF_TITLE_SIZE,
        leading=PDF_TITLE_SIZE + 4,
        textColor=HexColor(BRAND_COLOR),
    )
    body_style = ParagraphStyle(
        "body",
        fontName="Helvetica",
        fontSize=PDF_TEXT_SIZE,
        leading=PDF_LEADING,
        textColor=HexColor(TEXT_COLOR),
    )
    story: list[Any] = [
        Paragraph(escape(_clean(one_pager["title"])), title_style),
        Spacer(1, PDF_SPACE),
        Paragraph(escape(_clean(one_pager["summary"])), body_style),
        Spacer(1, PDF_SPACE),
        ListFlowable([Paragraph(escape(b), body_style) for b in _bullets(one_pager)]),
        Spacer(1, PDF_SPACE),
        Paragraph(escape(BRAND_NAME), body_style),
    ]
    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=LETTER, title=_clean(one_pager["title"]), author=BRAND_NAME, invariant=1
    ).build(story)
    return buffer.getvalue()


def _slide(deck: Any, title: str, lines: Sequence[str]) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    slide.shapes.title.text = title
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.size = Pt(DECK_TITLE_PT)
        run.font.color.rgb = RGBColor(*BRAND_RGB)  # type: ignore[no-untyped-call]
    if lines:
        box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(4))
        box.text_frame.word_wrap = True
        for index, line in enumerate(lines):
            paragraph = (
                box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
            )
            paragraph.text = line
            for run in paragraph.runs:
                run.font.size = Pt(DECK_TEXT_PT)


def _fixed_zip(blob: bytes) -> bytes:
    """Rewrite the PPTX zip with one timestamp on every entry, so equal decks are equal bytes."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(blob)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(item))
    return out.getvalue()


def render_deck_pptx(one_pager: Mapping[str, Any], landing_page: Mapping[str, Any]) -> bytes:
    """Render the sales deck: a title slide, the summary, the key points, and the call to
    action, as a PPTX."""
    deck = Presentation()
    _slide(deck, _clean(one_pager["title"]), [_clean(landing_page["subheadline"])])
    _slide(deck, "Summary", [_clean(one_pager["summary"])])
    _slide(deck, "Key points", _bullets(one_pager))
    _slide(deck, "Next step", [_clean(landing_page["call_to_action"])])
    deck.core_properties.author = BRAND_NAME
    deck.core_properties.created = FIXED_TIMESTAMP
    deck.core_properties.modified = FIXED_TIMESTAMP
    deck.core_properties.last_modified_by = BRAND_NAME
    buffer = io.BytesIO()
    deck.save(buffer)
    return _fixed_zip(buffer.getvalue())


def render_landing_page_html(landing_page: Mapping[str, Any]) -> str:
    """Render the landing page contract (headline, subheadline, body, call_to_action) as
    brand HTML. Every field is escaped."""
    parts = [p.strip() for p in _PARAGRAPH_BREAK.split(_clean(landing_page["body"])) if p.strip()]
    return _LANDING_TEMPLATE.format(
        headline=escape(_clean(landing_page["headline"])),
        subheadline=escape(_clean(landing_page["subheadline"])),
        paragraphs="\n".join(f"<p>{escape(p).replace(chr(10), '<br>')}</p>" for p in parts),
        cta=escape(_clean(landing_page["call_to_action"])),
        brand=escape(BRAND_NAME),
        color=BRAND_COLOR,
        text=TEXT_COLOR,
        font=FONT_STACK,
        width=PAGE_MAX_WIDTH,
    )


def render_graphic_svg(one_pager: Mapping[str, Any]) -> str:
    """Render the social graphic: the one-pager's title on a brand banner, from an SVG template."""
    title = " ".join(_clean(one_pager["title"]).split())
    if len(title) > SVG_TITLE_MAX_CHARS:
        title = title[: SVG_TITLE_MAX_CHARS - 1] + "…"
    return _SVG_TEMPLATE.format(
        w=SVG_WIDTH,
        h=SVG_HEIGHT,
        title=escape(title),
        brand=escape(BRAND_NAME),
        color=BRAND_COLOR,
        font=FONT_STACK,
        margin=SVG_MARGIN,
        title_y=SVG_TITLE_Y,
        brand_y=SVG_BRAND_Y,
        title_size=SVG_TITLE_SIZE,
        brand_size=SVG_BRAND_SIZE,
    )
